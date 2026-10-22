
import csv
import json
import queue
import time
import uuid
from datetime import datetime
from io import BytesIO, StringIO
from pathlib import Path

import pandas as pd
import streamlit as st
import paho.mqtt.client as mqtt
from openpyxl import Workbook, load_workbook
from PIL import Image

# =========================================================
# PlantFusion configuration
# =========================================================
MQTT_BROKER = "test.mosquitto.org"
MQTT_PORT = 1883
MQTT_TOPIC = "plantfusion/hima08656/sensors"

BASE_DIR = Path(__file__).resolve().parent
EXCEL_FILE = BASE_DIR / "PlantFusion_Data.xlsx"

HEADERS = [
    "timestamp",
    "air_temp_c",
    "air_humidity_pct",
    "soil_moisture_pct",
    "light_pct",
    "virtual_ph",
    "virtual_ec",
    "soil_temp_c",
]
SENSOR_COLUMNS = HEADERS[1:]

st.set_page_config(
    page_title="PlantFusion Dashboard",
    page_icon="🌿",
    layout="wide",
)

# =========================================================
# Excel storage helpers
# =========================================================
def open_or_create_excel():
    if EXCEL_FILE.exists():
        workbook = load_workbook(EXCEL_FILE)
        if "Sensor Data" in workbook.sheetnames:
            sheet = workbook["Sensor Data"]
        else:
            sheet = workbook.create_sheet("Sensor Data")
            sheet.append(HEADERS)
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sensor Data"
        sheet.append(HEADERS)

    if sheet.max_row == 1 and sheet.cell(1, 1).value is None:
        for col, header in enumerate(HEADERS, start=1):
            sheet.cell(row=1, column=col, value=header)

    return workbook, sheet


def excel_row_key(values):
    return tuple(
        "" if value is None else str(value).strip()
        for value in values[:len(HEADERS)]
    )


def append_readings_to_excel(readings):
    if not readings:
        return 0

    workbook = None
    try:
        workbook, sheet = open_or_create_excel()
        existing = set()

        for values in sheet.iter_rows(min_row=2, values_only=True):
            if values and any(value is not None for value in values):
                existing.add(excel_row_key(values))

        added = 0
        for data in readings:
            row = [data.get(key, "") for key in HEADERS]
            key = excel_row_key(row)

            if key not in existing:
                sheet.append(row)
                existing.add(key)
                added += 1

        workbook.save(EXCEL_FILE)
        return added
    finally:
        if workbook is not None:
            workbook.close()


def import_csv_text_to_excel(csv_text):
    reader = csv.DictReader(StringIO(csv_text))

    if not reader.fieldnames:
        return 0, "CSV file is empty."

    reader.fieldnames = [
        name.strip() if name else ""
        for name in reader.fieldnames
    ]

    missing = [
        name for name in HEADERS
        if name not in reader.fieldnames
    ]

    if missing:
        return 0, "Missing CSV columns: " + ", ".join(missing)

    workbook = None
    try:
        workbook, sheet = open_or_create_excel()
        existing = set()

        for values in sheet.iter_rows(min_row=2, values_only=True):
            if values and any(value is not None for value in values):
                existing.add(excel_row_key(values))

        added = 0
        for record in reader:
            row = [
                (record.get(key) or "").strip()
                for key in HEADERS
            ]

            if not any(row):
                continue

            key = excel_row_key(row)
            if key not in existing:
                sheet.append(row)
                existing.add(key)
                added += 1

        workbook.save(EXCEL_FILE)
        return added, None

    except Exception as error:
        return 0, str(error)

    finally:
        if workbook is not None:
            workbook.close()


def read_excel_dataframe():
    if not EXCEL_FILE.exists():
        return pd.DataFrame(columns=HEADERS)

    try:
        return pd.read_excel(
            EXCEL_FILE,
            sheet_name="Sensor Data"
        )
    except Exception:
        return pd.DataFrame(columns=HEADERS)


def get_excel_bytes():
    if not EXCEL_FILE.exists():
        workbook, _ = open_or_create_excel()
        workbook.save(EXCEL_FILE)
        workbook.close()

    return EXCEL_FILE.read_bytes()


# =========================================================
# MQTT connection
# =========================================================
@st.cache_resource
def start_mqtt():
    message_queue = queue.Queue()
    status = {"connected": False, "error": None}

    def on_connect(client, userdata, flags, reason_code, properties):
        if getattr(reason_code, "value", reason_code) == 0:
            status["connected"] = True
            status["error"] = None

            result, _ = client.subscribe(MQTT_TOPIC)
            if result != mqtt.MQTT_ERR_SUCCESS:
                status["error"] = "Could not subscribe to MQTT topic."
        else:
            status["connected"] = False
            status["error"] = (
                f"MQTT connection rejected: {reason_code}"
            )

    def on_disconnect(
        client, userdata, disconnect_flags, reason_code, properties
    ):
        status["connected"] = False

    def on_message(client, userdata, msg):
        try:
            data = json.loads(msg.payload.decode("utf-8"))

            if not isinstance(data, dict):
                return

            timestamp = (
                data.get("timestamp")
                or datetime.now().isoformat(timespec="seconds")
            )

            reading = {"timestamp": timestamp}

            for key in SENSOR_COLUMNS:
                reading[key] = data.get(key, "")

            message_queue.put(reading)

        except (json.JSONDecodeError, UnicodeDecodeError) as error:
            status["error"] = f"Invalid MQTT JSON: {error}"

    client = mqtt.Client(
        callback_api_version=mqtt.CallbackAPIVersion.VERSION2,
        client_id=f"plantfusion-dashboard-{uuid.uuid4().hex[:10]}",
        protocol=mqtt.MQTTv311,
    )

    client.on_connect = on_connect
    client.on_disconnect = on_disconnect
    client.on_message = on_message
    client.reconnect_delay_set(min_delay=1, max_delay=30)

    try:
        client.connect_async(
            MQTT_BROKER,
            MQTT_PORT,
            keepalive=60
        )
        client.loop_start()

    except Exception as error:
        status["error"] = str(error)

    return message_queue, status, client


message_queue, connection_status, mqtt_client = start_mqtt()

# =========================================================
# Process incoming MQTT readings
# =========================================================
if "sensor_history" not in st.session_state:
    st.session_state.sensor_history = []

new_readings = []

while True:
    try:
        reading = message_queue.get_nowait()
    except queue.Empty:
        break

    new_readings.append(reading)
    st.session_state.sensor_history.append(reading)

if new_readings:
    try:
        added = append_readings_to_excel(new_readings)

        if added:
            st.session_state.last_excel_save = (
                datetime.now().strftime("%H:%M:%S")
            )

    except Exception as error:
        st.session_state.excel_save_error = str(error)

st.session_state.sensor_history = (
    st.session_state.sensor_history[-300:]
)
history = st.session_state.sensor_history

# =========================================================
# Dashboard UI
# =========================================================
st.title("🌿 PlantFusion Monitoring Dashboard")
st.caption(
    "Wokwi → MQTT → Live Sensor Dashboard → Excel"
)

if connection_status["connected"]:
    st.success(
        f"MQTT connected — topic: {MQTT_TOPIC}"
    )
else:
    st.warning(
        "MQTT connection is pending or disconnected. "
        "Start Wokwi simulation and check the broker and topic."
    )

    if connection_status.get("error"):
        st.caption(
            f"Connection detail: {connection_status['error']}"
        )

if st.session_state.get("last_excel_save"):
    st.caption(
        "Last Excel save: "
        + st.session_state.last_excel_save
    )

if st.session_state.get("excel_save_error"):
    st.error(
        "Excel save error: "
        + st.session_state.excel_save_error
    )

if history:
    df = pd.DataFrame(history)

    for column in SENSOR_COLUMNS:
        if column in df.columns:
            df[column] = pd.to_numeric(
                df[column],
                errors="coerce"
            )

    latest = df.iloc[-1]

    st.subheader("📡 Latest Sensor Readings")

    c1, c2, c3, c4 = st.columns(4)

    c1.metric(
        "Air Temperature",
        f"{latest.get('air_temp_c', float('nan')):.1f} °C"
    )
    c2.metric(
        "Air Humidity",
        f"{latest.get('air_humidity_pct', float('nan')):.1f} %"
    )
    c3.metric(
        "Soil Moisture",
        f"{latest.get('soil_moisture_pct', float('nan')):.1f} %"
    )
    c4.metric(
        "Light Level",
        f"{latest.get('light_pct', float('nan')):.1f} %"
    )

    c5, c6, c7 = st.columns(3)

    c5.metric(
        "Virtual pH",
        f"{latest.get('virtual_ph', float('nan')):.2f}"
    )
    c6.metric(
        "Virtual EC",
        f"{latest.get('virtual_ec', float('nan')):.2f}"
    )
    c7.metric(
        "Soil Temperature",
        f"{latest.get('soil_temp_c', float('nan')):.1f} °C"
    )

    st.subheader("📈 Live Sensor Graphs")

    chart_df = df.copy()
    chart_df["timestamp"] = pd.to_datetime(
        chart_df["timestamp"],
        errors="coerce"
    )
    chart_df = chart_df.dropna(
        subset=["timestamp"]
    ).set_index("timestamp")

    left, right = st.columns(2)

    with left:
        st.markdown("**Air and Soil Temperature**")
        cols = [
            col for col in ["air_temp_c", "soil_temp_c"]
            if col in chart_df.columns
        ]
        if cols:
            st.line_chart(chart_df[cols])

    with right:
        st.markdown("**Soil Moisture and Light**")
        cols = [
            col for col in ["soil_moisture_pct", "light_pct"]
            if col in chart_df.columns
        ]
        if cols:
            st.line_chart(chart_df[cols])

    left2, right2 = st.columns(2)

    with left2:
        st.markdown("**Virtual pH**")
        st.line_chart(chart_df[["virtual_ph"]])

    with right2:
        st.markdown("**Virtual EC**")
        st.line_chart(chart_df[["virtual_ec"]])

    with st.expander("View recent live readings"):
        st.dataframe(
            df.iloc[::-1],
            use_container_width=True
        )

else:
    st.info(
        "Waiting for the first MQTT message. "
        "Start Wokwi simulation and publish JSON to "
        "plantfusion/hima08656/sensors."
    )

# =========================================================
# Import existing CSV data into Excel
# =========================================================
st.divider()
st.header("📊 Import Existing CSV Data")

st.write(
    "Upload your existing sensor CSV file. "
    "Its rows will be added to the Excel workbook; "
    "exact duplicate rows are skipped."
)

uploaded_csv = st.file_uploader(
    "Choose sensor CSV file",
    type=["csv"],
    key="pf_sensor_csv_upload",
)

if uploaded_csv is not None:
    try:
        csv_text = uploaded_csv.getvalue().decode("utf-8-sig")
        preview = pd.read_csv(StringIO(csv_text))

        st.caption(f"CSV rows found: {len(preview)}")
        st.dataframe(
            preview.head(10),
            use_container_width=True
        )

        if st.button(
            "Import CSV rows into Excel",
            key="pf_import_csv_button"
        ):
            added, error = import_csv_text_to_excel(csv_text)

            if error:
                st.error(error)
            else:
                st.success(
                    f"Imported {added} new row(s) into "
                    "PlantFusion_Data.xlsx"
                )
                st.rerun()

    except Exception as error:
        st.error(f"CSV import failed: {error}")

# =========================================================
# Excel preview and download
# =========================================================
st.divider()
st.header("📗 Excel Spreadsheet")

excel_df = read_excel_dataframe()
st.metric("Rows stored in Excel", len(excel_df))

if not excel_df.empty:
    st.dataframe(
        excel_df.tail(100).iloc[::-1],
        use_container_width=True
    )
else:
    st.info(
        "Excel file will receive rows when MQTT sensor "
        "messages arrive or when CSV is imported."
    )

st.download_button(
    "⬇️ Download latest PlantFusion_Data.xlsx",
    data=get_excel_bytes(),
    file_name="PlantFusion_Data.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    use_container_width=True,
)

# =========================================================
# Plant camera and image upload
# =========================================================
st.divider()
st.header("📷 Plant Camera")

camera_mode = st.radio(
    "Choose image source",
    [
        "Real Camera",
        "Virtual Sample Images",
        "Upload Plant Image"
    ],
    horizontal=True,
    key="pf_camera_mode",
)

if camera_mode == "Real Camera":
    st.caption(
        "Allow camera permission in your browser. "
        "Availability depends on your device and browser."
    )

    photo = st.camera_input(
        "Take a plant photo",
        key="pf_real_camera"
    )

    if photo is not None:
        st.image(
            photo,
            caption="Captured plant photo",
            use_container_width=True
        )

        st.download_button(
            "⬇️ Download captured photo",
            data=photo.getvalue(),
            file_name=f"plant_{datetime.now():%Y%m%d_%H%M%S}.jpg",
            mime="image/jpeg",
            key="pf_download_real_photo",
        )

elif camera_mode == "Virtual Sample Images":
    st.caption(
        "These are example images from the internet, "
        "not a live camera feed."
    )

    sample_urls = [
        (
            "Money Plant",
            "https://images.unsplash.com/photo-1497250681960-ef046c08a56e?auto=format&fit=crop&w=1000&q=85"
        ),
        (
            "Indoor Plant",
            "https://images.unsplash.com/photo-1501004318641-b39e6451bec6?auto=format&fit=crop&w=1000&q=85"
        ),
        (
            "Green Leaves",
            "https://images.unsplash.com/photo-1441974231531-c6227db76b6e?auto=format&fit=crop&w=1000&q=85"
        ),
    ]

    sample_name = st.selectbox(
        "Choose a sample",
        [item[0] for item in sample_urls]
    )
    sample_url = dict(sample_urls)[sample_name]

    try:
        import requests

        response = requests.get(
            sample_url,
            timeout=15
        )
        response.raise_for_status()

        sample_image = Image.open(
            BytesIO(response.content)
        ).convert("RGB")

        st.image(
            sample_image,
            caption=sample_name,
            use_container_width=True
        )

        buffer = BytesIO()
        sample_image.save(
            buffer,
            format="JPEG",
            quality=95
        )

        st.download_button(
            "⬇️ Download sample image",
            data=buffer.getvalue(),
            file_name=f"{sample_name.lower().replace(' ', '_')}.jpg",
            mime="image/jpeg",
            key="pf_download_sample",
        )

    except Exception as error:
        st.error(f"Could not load sample image: {error}")

else:
    uploaded_photo = st.file_uploader(
        "Upload a plant photo",
        type=["jpg", "jpeg", "png", "webp"],
        key="pf_uploaded_plant_photo",
    )

    if uploaded_photo is not None:
        try:
            image = Image.open(
                BytesIO(uploaded_photo.getvalue())
            ).convert("RGB")

            st.image(
                image,
                caption=f"Uploaded: {uploaded_photo.name}",
                use_container_width=True
            )

            st.download_button(
                "⬇️ Download uploaded plant photo",
                data=uploaded_photo.getvalue(),
                file_name=uploaded_photo.name,
                mime=uploaded_photo.type or "image/jpeg",
                key="pf_download_uploaded_photo",
            )

        except Exception as error:
            st.error(f"Could not open image: {error}")

# =========================================================
# Refresh dashboard
# =========================================================
st.caption(
    "Live readings refresh approximately every 2 seconds "
    "while this dashboard is open."
)

time.sleep(2)
st.rerun()
